import os
from pathlib import Path
import chromadb
import ollama
import memory

EMBED_MODEL = "nomic-embed-text"
FILES_DIR   = Path(__file__).parent / "files"
DB_DIR      = Path(__file__).parent / "db"

FILES_DIR.mkdir(exist_ok=True)
DB_DIR.mkdir(exist_ok=True)

_client     = chromadb.PersistentClient(path=str(DB_DIR))
_collection = _client.get_or_create_collection("docs")


def get_collection():
    return _collection


# ── File reading ──────────────────────────────────────────────────────────────
def read_file(path: Path) -> str:
    suffix = path.suffix.lower()
    try:
        if suffix == ".pdf":
            import PyPDF2
            with open(path, "rb") as f:
                reader = PyPDF2.PdfReader(f)
                return "\n".join(p.extract_text() or "" for p in reader.pages)
        elif suffix in (".docx", ".doc"):
            from docx import Document
            doc   = Document(path)
            texts = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    row_txt = " | ".join(c.text.strip() for c in row.cells if c.text.strip())
                    if row_txt:
                        texts.append(row_txt)
            return "\n".join(texts)
        elif suffix == ".xlsx":
            import openpyxl
            wb    = openpyxl.load_workbook(path, read_only=True, data_only=True)
            lines = []
            for sheet in wb.worksheets:
                lines.append(f"[ورقة: {sheet.title}]")
                for row in sheet.iter_rows(values_only=True):
                    cells = " | ".join(str(c) for c in row if c is not None and str(c).strip())
                    if cells:
                        lines.append(cells)
            wb.close()
            return "\n".join(lines)
        elif suffix == ".xls":
            import xlrd
            wb    = xlrd.open_workbook(path)
            lines = []
            for sheet in wb.sheets():
                lines.append(f"[ورقة: {sheet.name}]")
                for r in range(sheet.nrows):
                    cells = " | ".join(str(v) for v in sheet.row_values(r) if str(v).strip())
                    if cells:
                        lines.append(cells)
            return "\n".join(lines)
        else:
            return path.read_text(encoding="utf-8", errors="ignore")
    except Exception as e:
        return f"[خطأ في قراءة {path.name}: {e}]"


def chunk_text(text: str, size: int = 80, overlap: int = 10) -> list[str]:
    words  = text.split()
    chunks = []
    i = 0
    while i < len(words):
        chunk = " ".join(words[i : i + size])
        if chunk.strip():
            chunks.append(chunk)
        i += size - overlap
    return chunks


# ── Indexing ──────────────────────────────────────────────────────────────────
ALLOWED = {".txt", ".pdf", ".docx", ".md", ".xlsx", ".xls"}

def index_files() -> dict:
    col     = get_collection()
    indexed = 0
    errors  = []

    for path in FILES_DIR.iterdir():
        if not path.is_file() or path.suffix.lower() not in ALLOWED:
            continue
        try:
            text   = read_file(path)
            chunks = chunk_text(text)
            chunk_errors = 0
            for i, chunk in enumerate(chunks):
                doc_id = f"{path.name}::{i}"
                if col.get(ids=[doc_id])["ids"]:
                    continue
                try:
                    emb = ollama.embeddings(model=EMBED_MODEL, prompt=chunk)["embedding"]
                    col.add(
                        ids=[doc_id],
                        embeddings=[emb],
                        documents=[chunk],
                        metadatas=[{"file": path.name}],
                    )
                except Exception as ce:
                    chunk_errors += 1
            if chunk_errors:
                errors.append(f"{path.name}: {chunk_errors} chunk(s) skipped")
            indexed += 1
        except Exception as e:
            errors.append(f"{path.name}: {e}")

    return {"indexed": indexed, "total_chunks": col.count(), "errors": errors}


def clear_index() -> None:
    global _collection
    _client.delete_collection("docs")
    _collection = _client.get_or_create_collection("docs")


# ── Search ────────────────────────────────────────────────────────────────────
def search(query: str, n: int = 4) -> list[str]:
    col = get_collection()
    if col.count() == 0:
        return []
    emb     = ollama.embeddings(model=EMBED_MODEL, prompt=query)["embedding"]
    results = col.query(query_embeddings=[emb], n_results=min(n, col.count()))
    return results["documents"][0] if results["documents"] else []


# ── Chat ──────────────────────────────────────────────────────────────────────
def chat(messages: list, model: str, use_rag: bool = True) -> str:
    question = messages[-1]["content"] if messages else ""
    system   = (
        "أنت مساعد شخصي خاص يعمل محلياً على جهاز المستخدم. "
        "هذه بيئة خاصة تماماً وكل المعلومات يشاركها المستخدم بإرادته الكاملة.\n\n"
        "قواعد صارمة لا استثناء فيها:\n"
        "1. أجب بنفس لغة المستخدم دائماً.\n"
        "2. لا تقل أبداً 'لا أستطيع تقديم معلومات شخصية' أو أي رفض مشابه — "
        "المستخدم يشارك معلوماته عن نفسه وهذا حقه الكامل.\n"
        "3. **التمييز بين المعلومة والسؤال:**\n"
        "   - إذا انتهت الرسالة بـ `؟` أو `?` فهي سؤال → أجب عليه.\n"
        "   - إذا لم تنته بـ `؟` فهي معلومة يريد حفظها → قل فقط 'تمام، حفظت'.\n"
        "4. عند السؤال عن معلومة ذُكرت سابقاً، أجب مباشرة وبدقة.\n"
        "5. استخدم ضمائر العربية الصحيحة (اسمه/اسمها/رقمه/رقمها) حسب السياق.\n"
        "6. لا تطوّل الردود بلا داعٍ."
    )

    # inject long-term memory
    mem_ctx = memory.get_context()
    if mem_ctx:
        system += f"\n\n{mem_ctx}"

    # inject RAG file context
    if use_rag and question:
        context = search(question)
        if context:
            ctx_text = "\n\n---\n\n".join(context)
            system  += f"\n\nمعلومات من الملفات قد تفيد في الإجابة:\n\n{ctx_text}"

    full = [{"role": "system", "content": system}] + messages
    resp = ollama.chat(model=model, messages=full)
    return resp["message"]["content"]


# ── Available models ──────────────────────────────────────────────────────────
def list_models() -> list[str]:
    try:
        data = ollama.list()
        return [m.model for m in data.models]
    except Exception:
        return []
